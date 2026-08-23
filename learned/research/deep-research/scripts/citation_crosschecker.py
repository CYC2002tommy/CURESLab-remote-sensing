import sys
import re
try:
    import docx
except ImportError:
    print("Please install python-docx: pip install python-docx")
    sys.exit(1)

def main(docx_path):
    doc = docx.Document(docx_path)
    main_body, refs_body = [], []
    in_refs = False

    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.upper() == 'REFERENCES':
            in_refs = True
            continue
        if in_refs and txt:
            refs_body.append(txt)
        elif not in_refs:
            main_body.append(txt)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                main_body.append(cell.text)

    main_text = " ".join(main_body)
    
    print(f"Found {len(refs_body)} references in bibliography.")
    
    # 1. Check for unused references (In Bib, NOT in text)
    unused = []
    for ref in refs_body:
        year_match = re.search(r'\(\d{4}[a-z]?\)', ref)
        if not year_match: continue
        year = year_match.group(0).strip('()')
        first_word = ref.split(',')[0].strip()
        if first_word.lower() not in main_text.lower() or year not in main_text:
            unused.append(ref)
            
    print("\n--- UNUSED REFERENCES (In Bib, NOT in text) ---")
    for r in unused: print(f"- {r[:60]}...")
        
    # 2. Check for missing references (In text, NOT in Bib)
    cites = []
    for m in re.findall(r'\(([^)]*\d{4}[a-z]?)\)', main_text):
        for sub in m.split(';'):
            if re.search(r'\d{4}', sub): cites.append(sub.strip())
    for m in re.findall(r'([A-Z][A-Za-z\-]+(?:\s+et\s+al\.?)?)\s+\((\d{4}[a-z]?)\)', main_text):
        cites.append(f"{m[0]} {m[1]}")
        
    missing = []
    for cite in set(cites):
        year_match = re.search(r'(\d{4})', cite)
        if not year_match: continue
        year = year_match.group(1)
        names = re.findall(r'[A-Z][A-Za-z\-]+', cite)
        if not names: continue
        
        found = False
        for ref in refs_body:
            if year in ref and any(n.lower() in ref.lower() for n in names):
                found = True; break
        if not found: missing.append(cite)
        
    print("\n--- MISSING REFERENCES (In text, NOT in Bib) ---")
    for m in missing: print(f"- {m}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 citation_crosschecker.py <file.docx>")
        sys.exit(1)
    main(sys.argv[1])